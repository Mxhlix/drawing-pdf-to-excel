"""転記先の Excel の様式（5つ）と、設定.xlsx（部屋名の一覧と区分）を作る。

実際の様式は使わず、よくある形で作った見本。実際の様式に合わせるときは、
config/*.yaml の欄の位置（cell・start_row・columns の並び）を合わせるだけで差し替えられる。

  $ python make_templates.py
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from zumen.mapper import load_configs

HERE = Path(__file__).parent
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DDE5EE")
FONT = "游ゴシック"


def template(cfg):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg["sheet"]
    ws["A1"] = cfg["title"]
    ws["A1"].font = Font(name=FONT, size=16, bold=True)
    for f in cfg["header"]:
        c = ws[f["cell"]]
        lab = ws.cell(row=c.row, column=c.column - 1, value=f["field"])
        lab.font = Font(name=FONT, size=10, bold=True)
        lab.fill = HEAD
        lab.border = BOX
        c.border = BOX
        c.font = Font(name=FONT, size=11)
    row = cfg["start_row"] - 1
    cols = ["No."] + cfg["records"]["columns"]
    for k, name in enumerate(cols, start=1):
        h = ws.cell(row=row, column=k, value=name)
        h.font = Font(name=FONT, size=10, bold=True)
        h.fill = HEAD
        h.border = BOX
        h.alignment = Alignment(horizontal="center")
        for r in range(row + 1, row + 21):
            ws.cell(row=r, column=k).border = BOX
            ws.cell(row=r, column=k).font = Font(name=FONT, size=11)
    ws.column_dimensions["A"].width = 6
    for k in range(2, len(cols) + 4):
        ws.column_dimensions[openpyxl.utils.get_column_letter(k)].width = 18
    ws.column_dimensions["C"].width = 26
    wb.save(HERE / "templates" / cfg["template"])


ROOMS = [("LDK", "居室"), ("リビング", "居室"), ("和室", "居室"), ("洋室", "居室"), ("寝室", "居室"),
         ("子供室", "居室"), ("押入", "収納"), ("収納", "収納"), ("WIC", "収納"), ("パントリー", "収納"),
         ("玄関", "その他"), ("ホール", "その他"), ("廊下", "その他"), ("階段", "その他"),
         ("トイレ", "水回り"), ("洗面脱衣室", "水回り"), ("浴室", "水回り"), ("キッチン", "水回り"),
         ("事務室", "執務"), ("会議室", "執務"), ("応接室", "執務"), ("給湯室", "水回り"), ("更衣室", "その他"),
         ("倉庫", "収納"), ("サーバー室", "その他"), ("待合室", "共用"), ("受付", "共用"), ("診察室1", "診療"),
         ("診察室2", "診療"), ("処置室", "診療"), ("スタッフ室", "その他"), ("売場", "店舗"),
         ("レジカウンター", "店舗"), ("通路", "共用"), ("バックヤード", "収納"), ("事務所", "執務"), ("休憩室", "その他")]


EQUIP = ["天井カセット形エアコン", "天井吊形エアコン", "壁掛形エアコン", "床置形エアコン", "全熱交換器",
         "天井埋込換気扇", "パイプファン", "中間ダクトファン", "有圧換気扇"]


FINISH = ["長尺塩ビシート", "塩ビタイル", "磁器質タイル", "タイルカーペット", "フローリング", "防塵塗装",
          "ビニル幅木", "ソフト幅木", "木製幅木", "ビニルクロス", "塗装", "岩綿吸音板", "ケイカル板", "化粧石膏ボード",
          "直天井", "なし"]


def settings():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "部屋名"
    ws.append(["部屋名", "区分"])
    for r in ROOMS:
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True)
        c.fill = HEAD
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws2 = wb.create_sheet("機器名称")
    ws2.append(["名称"])
    for name in EQUIP:
        ws2.append([name])
    ws2["A1"].font = Font(name=FONT, bold=True)
    ws2["A1"].fill = HEAD
    ws2.column_dimensions["A"].width = 28
    ws3 = wb.create_sheet("仕上材")
    ws3.append(["材料"])
    for name in FINISH:
        ws3.append([name])
    ws3["A1"].font = Font(name=FONT, bold=True)
    ws3["A1"].fill = HEAD
    ws3.column_dimensions["A"].width = 24
    wb.save(HERE / "設定.xlsx")


if __name__ == "__main__":
    (HERE / "templates").mkdir(exist_ok=True)
    settings()
    for cfg in load_configs().values():
        template(cfg)
        print(cfg["template"])
