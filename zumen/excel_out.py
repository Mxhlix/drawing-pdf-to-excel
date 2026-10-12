"""振り分けた値を、様式の Excel に書き込む。確かめが済んでいない欄は黄色にしてメモを付ける。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

HERE = Path(__file__).resolve().parent.parent
YELLOW = PatternFill("solid", fgColor="FFF2A8")


def _num(v: str):
    try:
        return int(v) if v.isdigit() else float(v) if v.replace(".", "", 1).isdigit() else v
    except ValueError:
        return v


def write(result: dict, cfg: dict, pdf_name: str) -> Path:
    wb = openpyxl.load_workbook(HERE / "templates" / cfg["template"])
    ws = wb[cfg["sheet"]]

    def put(ref_or_rc, c):
        cellobj = ws[ref_or_rc] if isinstance(ref_or_rc, str) else ws.cell(*ref_or_rc)
        cellobj.value = _num(c["value"]) if c["value"] else None
        if c.get("flag"):
            cellobj.fill = YELLOW
            cellobj.comment = Comment("要確認: " + " / ".join(c["reasons"]), "転記ツール")

    for f in cfg["header"]:
        put(f["cell"], result["header"][f["field"]])
    row = cfg["start_row"]
    for n, rec in enumerate(result["records"], start=1):
        ws.cell(row=row, column=1, value=n)
        for k, col in enumerate(cfg["records"]["columns"], start=2):
            if col in rec:
                put((row, k), rec[col])
        row += 1
    out = HERE / "output" / f"{Path(pdf_name).stem}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    return out
