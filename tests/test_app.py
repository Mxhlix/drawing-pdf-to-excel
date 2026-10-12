"""画面の裏の API を、文字ありPDFで通しで確かめる（読む → 様式を変える → Excel に書き出す → 取り出す）。"""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

import app as app_module
from evaluate import HERE


@pytest.fixture(scope="module")
def client():
    return TestClient(app_module.app)


def read(client, name):
    with open(HERE / "samples" / name, "rb") as f:
        r = client.post("/api/read", files={"file": (name, f, "application/pdf")})
    assert r.status_code == 200
    return r.json()


def test_configs_lists_five_forms(client):
    keys = {c["key"] for c in client.get("/api/configs").json()}
    assert keys == {"house", "aircon", "lighting", "interior", "part"}


def test_read_edit_export_download(client):
    d = read(client, "house_text.pdf")
    assert d["kind"] == "text" and d["key"] == "house"
    assert d["header"]["図番"]["value"] == "A-03"
    rooms = {r["部屋名"]["value"] for r in d["records"]}
    assert "LDK" in rooms

    # 担当者が画面で1欄直して、確かめの印を付けたまま書き出す
    header = dict(d["header"])
    header["工事名称"] = {**header["工事名称"], "value": "直した工事名", "flag": True, "reasons": ["手で直した"]}
    r = client.post(f"/api/export/{d['id']}", json={"header": header, "records": d["records"]})
    assert r.status_code == 200
    out = r.json()

    got = client.get(f"/api/download/{out['file']}")
    assert got.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(got.content))["部屋一覧"]
    assert ws["C3"].value == "直した工事名"
    assert ws["C3"].fill.fgColor.rgb.endswith("FFF2A8")      # 印の付いた欄は黄色
    assert "手で直した" in ws["C3"].comment.text
    assert ws["F3"].value == "A-03"
    assert ws.cell(row=7, column=1).value == 1               # 7行目から番号付きで並ぶ
    assert ws.cell(row=7 + len(d["records"]) - 1, column=2).value is not None
    (HERE / "output" / out["file"]).unlink()


def test_remap_switches_form(client):
    d = read(client, "aircon_text.pdf")
    assert d["key"] == "aircon"
    r = client.post(f"/api/remap/{d['id']}", json={"key": "part"})
    assert r.status_code == 200 and r.json()["key"] == "part"


@pytest.mark.parametrize("name", ["..%5Capp.py", "..%5C..%5Capp.py", "..", "missing.xlsx"])
def test_download_stays_inside_output_folder(client, name):
    r = client.get(f"/api/download/{name}")
    assert r.status_code == 404
